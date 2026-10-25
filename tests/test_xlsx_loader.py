# -*- coding: utf-8 -*-
"""v0.9.120.1 — Excel (xlsx) ingest + upload UX.

User report: 「chat ui 上傳 excel 無法正常上傳」. The upload endpoint
itself always accepted .xlsx (only dangerous executables are refused)
— the file just became a DEAD END afterwards and the only upload
entry points were the 📎 button and image-only paste:

* the ingest pipeline skip-listed .xlsx (nothing could parse it)
* read_file treats .xlsx as binary
* onChatPaste collected ONLY image/* clipboard items (a copied Excel
  file from Explorer did nothing)
* there was no drag-and-drop upload at all

Covers: the openpyxl sheet→pipe-table loader (sheets/CJK/empty
rows/cell caps), the pipeline skip-set change, and the two JS upload
entry points (paste-any-file + drag-and-drop).
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


# ── loader ────────────────────────────────────────────────────────
class TestXlsxLoader:
    @pytest.fixture()
    def workbook_path(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "營收表"
        ws.append(["月份", "營收", "備註"])
        ws.append(["2026-01", 123456, ""])
        ws.append([])              # fully empty row → skipped
        ws.append([None, None])    # all-None row → skipped
        ws2 = wb.create_sheet("長文本")
        ws2["A1"] = "字" * 250     # over the 200-char cell cap
        wb.create_sheet("空表")    # no rows → no section at all
        p = tmp_path / "t.xlsx"
        wb.save(p)
        return p

    def test_load_file_routes_xlsx(self, workbook_path):
        from src.ingest.loader import load_file

        text = load_file(workbook_path)
        assert text, "xlsx must load (was a dead end before 120.1)"

    def test_sheet_sections_and_rows(self, workbook_path):
        from src.ingest.loader import load_file

        text = load_file(workbook_path)
        assert "## Sheet: 營收表" in text
        assert "月份 | 營收 | 備註" in text
        assert "2026-01 | 123456" in text
        # the all-empty sheet produces NO section
        assert "## Sheet: 空表" not in text

    def test_cell_cap_truncation_marker(self, workbook_path):
        from src.ingest.loader import load_file

        text = load_file(workbook_path)
        assert ("字" * 200 + "…") in text
        assert "字" * 201 not in text

    def test_fileloader_class_same_output(self, workbook_path):
        from src.ingest.loader import FileLoader, load_file

        assert FileLoader.load_file(workbook_path) == \
            load_file(workbook_path)

    def test_unknown_file_returns_empty(self, tmp_path):
        from src.ingest.loader import load_file

        p = tmp_path / "broken.xlsx"
        p.write_bytes(b"PK\x03\x04 not a real workbook")
        assert load_file(p) == ""        # loader failure → "" (graceful)


# ── pipeline skip-set ─────────────────────────────────────────────
class TestPipelineSkipSet:
    def test_xlsx_ingestible_xls_xlsm_still_skipped(self):
        from src.ingest.pipeline import _SKIP_EXTS

        assert ".xlsx" not in _SKIP_EXTS
        assert ".xls" in _SKIP_EXTS and ".xlsm" in _SKIP_EXTS


# ── JS upload entry points ───────────────────────────────────────
class TestJsUploadEntries:
    @staticmethod
    def _js():
        return (Path(__file__).resolve().parent.parent
                / "gui/assets/app.js").read_text(encoding="utf-8")

    def test_paste_accepts_non_image_files(self):
        js = self._js()
        assert "kind !== 'file') continue;" in js
        assert "uploadFiles(files);" in js
        # the image branch keeps the v0.9.29.0 contract
        assert "handlePastedImage(f, idx)" in js

    def test_drag_and_drop_wired(self):
        js = self._js()
        assert "function bindDrop()" in js
        assert "bindDrop();" in js            # wired inside bindEvents
        assert "uploadFiles(dropped);" in js
        assert "body.classList.add('drag-over')" in js

    def test_style_css_has_drag_over_cue(self):
        css = (Path(__file__).resolve().parent.parent
               / "gui/assets/style.css").read_text(encoding="utf-8")
        assert "body.drag-over::after" in css


# ── dependency contract ──────────────────────────────────────────
class TestDependencyContract:
    def test_openpyxl_declared(self):
        pyproject = (Path(__file__).resolve().parent.parent
                     / "pyproject.toml").read_text(encoding="utf-8")
        assert "openpyxl>=3.1" in pyproject